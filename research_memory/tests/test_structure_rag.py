from __future__ import annotations

from pathlib import Path

from research_memory.pipeline.structure_rag import (
    StructureBlock,
    build_chunks_for_strategy,
    _group_sections,
)


def test_group_sections_flushes_on_heading_and_page() -> None:
    blocks = [
        StructureBlock("1. 개요", "heading", page=1, location="heading 1"),
        StructureBlock("본문A", "paragraph", page=1, location="문단 1"),
        StructureBlock("본문B", "paragraph", page=2, location="문단 2"),
        StructureBlock("표내용", "table", page=2, location="표 1"),
    ]
    sections = _group_sections(blocks)
    assert len(sections) == 3
    assert sections[0]["heading_path"] == "1. 개요"
    assert sections[0]["text"] == "본문A"
    assert sections[1]["page"] == 2
    assert sections[2]["block_type"] == "table"


def test_build_chunks_tag_strategy(tmp_path: Path) -> None:
    # Minimal fake docx via python-docx
    from docx import Document

    path = tmp_path / "sample.docx"
    doc = Document()
    doc.add_heading("평가 기준", level=1)
    doc.add_paragraph("정확도는 상대 오차율의 평균으로 산출한다. " * 5)
    doc.save(path)

    legacy = build_chunks_for_strategy(path, strategy="legacy", document_id="d1")
    structured = build_chunks_for_strategy(path, strategy="structure", document_id="d1")
    assert legacy and all(c.chunk_strategy == "legacy" for c in legacy)
    assert structured and all(c.chunk_strategy == "structure" for c in structured)
    assert any(c.heading_path for c in structured)
